#!/usr/bin/env python3
"""Testes de fumaça do pipeline.

Rodam sem rede, sem credenciais e sem cota de API. Os testes marcados como
*ffmpeg* são pulados automaticamente se o binário não estiver no PATH.

    pytest -q test_pipeline.py          # com pytest
    python test_pipeline.py             # standalone, sem dependências extras
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

import prompts_data
from config import load_settings
from generator import ASPECT_FALLBACKS, VeoGenerator, _looks_unsupported
from processor import build_loop_filter, sha256_file
from tracker import COLUMNS, REQUIRED_COLUMNS, RowUpdate, STATUS_DONE
from utils import is_retryable, retry

HAS_FFMPEG = shutil.which("ffmpeg") is not None and shutil.which("ffprobe") is not None


# --------------------------------------------------------------------------- #
# Matriz de prompts
# --------------------------------------------------------------------------- #
def test_matriz_tem_20_clipes():
    assert prompts_data.matrix_size() == 20
    assert len(list(prompts_data.iter_matrix())) == 20
    assert len(prompts_data.ACTOR_ORDER) == 4
    assert len(prompts_data.STATE_ORDER) == 5


def test_todo_prompt_carrega_as_regras_clinicas():
    """Boca fechada, câmera fixa e ausência de sinal patognomônico em 20/20."""
    for actor_id, state_id in prompts_data.iter_matrix():
        prompt = prompts_data.build_prompt(actor_id, state_id)
        low = prompt.lower()
        assert "mouth stays completely closed" in low, f"{actor_id}/{state_id}: boca"
        assert "no lip-sync" in low, f"{actor_id}/{state_id}: lip-sync"
        assert "no camera movement whatsoever" in low, f"{actor_id}/{state_id}: câmera"
        assert "never a diagnosis-revealing sign" in low, f"{actor_id}/{state_id}: patognomônico"
        assert len(prompt) > 800, f"{actor_id}/{state_id}: prompt curto demais"


def test_prompts_sao_todos_distintos():
    prompts = {
        prompts_data.build_prompt(a, s) for a, s in prompts_data.iter_matrix()
    }
    assert len(prompts) == 20, "há prompts duplicados na matriz"


def test_negative_prompt_bloqueia_fala_e_camera():
    for termo in ("talking", "open mouth", "lip sync", "camera movement", "zoom", "jaundice"):
        assert termo in prompts_data.NEGATIVE_PROMPT


def test_direcoes_de_estado_nao_presumem_postura():
    """A03 está deitada em leito a 30°.

    Uma direção compartilhada que diga "sentado" ou "na cadeira" contradiz o
    SETTING dessa persona dentro do próprio prompt — exatamente o tipo de
    conflito que faz o modelo escolher sozinho qual instrução ignorar.
    """
    proibidos = ("chair", "sits ", "sitting", "seated", "stands", "standing")
    for state_id, state in prompts_data.STATES.items():
        texto = f"{state['direcao']} {state['respiracao']}".lower()
        for termo in proibidos:
            assert termo not in texto, (
                f"estado {state_id!r} presume postura ({termo!r}); "
                "quem define postura é o SETTING de cada persona"
            )


def test_a03_nao_tem_contradicao_postural():
    prompt = prompts_data.build_prompt("A03", "dispneia").lower()
    assert "lying in a hospital ward bed" in prompt
    assert "in the chair" not in prompt
    assert "sits upright" not in prompt


def test_persona_deitada_tem_override_postural():
    """A03 está em leito a 30°: direções de 'sentar/inclinar' não se aplicam."""
    assert prompts_data.ACTORS["A03"].get("postura_override")
    assert "POSTURE NOTE" in prompts_data.build_prompt("A03", "dispneia")
    assert "POSTURE NOTE" not in prompts_data.build_prompt("A01", "dispneia")


def test_ator_ou_estado_desconhecido_falha_claro():
    for fn, arg in ((prompts_data.get_actor, "A99"), (prompts_data.get_state, "euforia")):
        try:
            fn(arg)
        except KeyError as exc:
            assert "Disponíveis" in str(exc)
        else:  # pragma: no cover
            raise AssertionError(f"{fn.__name__} deveria ter falhado para {arg!r}")


# --------------------------------------------------------------------------- #
# Filtergraph do loop
# --------------------------------------------------------------------------- #
def test_filtergraph_implementa_a_geometria_especificada():
    """A=[2,4) ⊕ B=[0,2) com xfade de 0,5s no offset 1,5s ⇒ 3,5s de saída."""
    graph = build_loop_filter(segment=2.0, crossfade=0.5, size=720, fps=24)
    assert "trim=start=2.0000:end=4.0000" in graph      # segmento A
    assert "trim=start=0:end=2.0000" in graph           # segmento B
    assert "xfade=transition=fade:duration=0.5000:offset=1.5000" in graph
    assert "crop=720:720" in graph                      # 1:1 exato
    assert "fps=24" in graph
    assert "format=yuv420p" in graph
    # Ordem importa: A precisa ser declarado antes de B para o loop fechar.
    assert graph.index("start=2.0000") < graph.index("start=0:end=2.0000")


def test_offset_do_xfade_e_sempre_segmento_menos_crossfade():
    """offset errado = corte visível na emenda; é a aritmética crítica do loop."""
    import re

    for segment, crossfade in ((2.0, 0.5), (3.0, 1.0), (1.5, 0.25)):
        graph = build_loop_filter(segment=segment, crossfade=crossfade, size=720, fps=24)
        offset = float(re.search(r"offset=([0-9.]+)", graph).group(1))
        assert abs(offset - (segment - crossfade)) < 1e-6
        # Duração resultante = offset + crossfade + (segmento - crossfade)
        assert abs((offset + segment) - (segment * 2 - crossfade)) < 1e-6


# --------------------------------------------------------------------------- #
# Retentativas
# --------------------------------------------------------------------------- #
def test_rate_limit_e_classificado_como_transitorio():
    assert is_retryable(RuntimeError("429 RESOURCE_EXHAUSTED: quota exceeded"))
    assert is_retryable(RuntimeError("503 Service Unavailable"))
    assert is_retryable(ConnectionError("connection reset by peer"))
    assert not is_retryable(ValueError("prompt inválido"))


def test_retry_persiste_ate_o_sucesso():
    tentativas = {"n": 0}

    def flaky():
        tentativas["n"] += 1
        if tentativas["n"] < 3:
            raise RuntimeError("429 rate limit")
        return "ok"

    assert retry(flaky, what="teste", attempts=5, base_delay=0.001, max_delay=0.01) == "ok"
    assert tentativas["n"] == 3


def test_retry_nao_insiste_em_erro_permanente():
    tentativas = {"n": 0}

    def sempre_invalido():
        tentativas["n"] += 1
        raise ValueError("argumento inválido")

    try:
        retry(sempre_invalido, what="teste", attempts=5, base_delay=0.001)
    except ValueError:
        pass
    assert tentativas["n"] == 1, "erro permanente não deve ser repetido"


def test_fallback_de_parametros_do_veo():
    assert _looks_unsupported(RuntimeError("INVALID_ARGUMENT: aspect_ratio must be one of 16:9"))
    assert not _looks_unsupported(RuntimeError("429 rate limit"))
    assert "16:9" in ASPECT_FALLBACKS


def test_generator_percorre_a_escada_de_fallback():
    """Recusa de 1:1 pela API deve cair para 16:9 sem abortar o clipe."""
    try:
        from google.genai import types  # type: ignore
    except ImportError:  # pragma: no cover
        print("  (pulado: google-genai não instalado)")
        return

    settings = load_settings({"google_api_key": "x", "veo_aspect_ratio": "1:1"})
    generator = VeoGenerator(settings)
    chamadas = []

    class FakeVideo:
        video_bytes = b"\x00" * 64

    class FakeOp:
        done = True
        error = None

        class response:  # noqa: N801
            generated_videos = [type("G", (), {"video": FakeVideo()})()]

    class FakeModels:
        def generate_videos(self, model, prompt, config):
            chamadas.append(config.aspect_ratio)
            if config.aspect_ratio == "1:1":
                raise RuntimeError("INVALID_ARGUMENT: aspect_ratio must be one of [16:9, 9:16]")
            return FakeOp()

    class FakeClient:
        models = FakeModels()

    generator._client = FakeClient()
    generator._types = types

    with tempfile.TemporaryDirectory() as tmp:
        dst = Path(tmp) / "raw.mp4"
        result = generator.generate("prompt", "negativo", dst, label="teste")
        assert dst.is_file() and dst.stat().st_size == 64, "o MP4 bruto deve ter sido escrito"

    assert chamadas[0] == "1:1", "deveria ter tentado o valor pedido primeiro"
    assert result.aspect_ratio == "16:9", "deveria ter caído para 16:9"
    # Um aspect ratio recusado não deve ser reensaiado com outras durações:
    # exatamente uma chamada desperdiçada, não três.
    assert chamadas.count("1:1") == 1, f"cota desperdiçada: {chamadas}"
    assert len(chamadas) == 2, f"esperava 1 recusa + 1 sucesso, obtive {chamadas}"
    # A combinação aceita fica memorizada para os 19 clipes seguintes.
    assert generator._aspect_ratio == "16:9"


# --------------------------------------------------------------------------- #
# Modo manual (vídeos gerados no app do Gemini/Flow)
# --------------------------------------------------------------------------- #
def test_prompt_manual_e_um_bloco_unico_com_exclusoes():
    """O app não tem campo de prompt negativo: as exclusões vão no próprio texto."""
    from run_pipeline import manual_prompt

    texto = manual_prompt("A02", "dor_intensa")
    assert texto.startswith(prompts_data.build_prompt("A02", "dor_intensa"))
    assert "AVOID: " in texto and "open mouth" in texto


def test_exportacao_gera_um_prompt_por_clipe_e_checklist():
    from run_pipeline import export_prompts

    with tempfile.TemporaryDirectory() as tmp:
        settings = load_settings({"work_dir": Path(tmp) / "build", "actors": ["A03"]})
        index = export_prompts(settings, Path(tmp) / "prompts")
        arquivos = sorted(p.name for p in (Path(tmp) / "prompts").glob("A03_*.txt"))
        assert len(arquivos) == 5
        checklist = index.read_text(encoding="utf-8")
        assert "A03_basal_raw.mp4" in checklist and "--manual-raw" in checklist


def test_saida_de_dry_run_nunca_passa_por_concluida_em_execucao_real():
    from run_pipeline import already_done, prompt_fingerprint

    prompt = prompts_data.build_prompt("A01", "basal")
    with tempfile.TemporaryDirectory() as tmp:
        settings = load_settings({"work_dir": Path(tmp)})
        out = settings.actor_out_dir("A01")
        out.mkdir(parents=True)
        (out / "basal.mp4").write_bytes(b"x")
        (out / "basal.jpg").write_bytes(b"x")
        manifest = {"clips": {"basal": {
            "prompt_sha256": prompt_fingerprint(prompt),
            "generation": {"model": "(dry-run)"},
            "video": {"file": "basal.mp4"},
            "poster": {"file": "basal.jpg"},
        }}}
        assert already_done(settings, manifest, "A01", "basal", prompt) is None
        settings.dry_run = True
        assert already_done(settings, manifest, "A01", "basal", prompt) is not None


def test_modo_manual_dispensa_a_chave_da_api():
    settings = load_settings({"google_api_key": "", "manual_raw": True, "drive_enabled": False})
    assert not any("GOOGLE_API_KEY" in p for p in settings.validate())


# --------------------------------------------------------------------------- #
# Planilha
# --------------------------------------------------------------------------- #
def test_colunas_obrigatorias_presentes_e_na_ordem():
    esperado = [
        "Ator", "Estado", "Status", "Tamanho MP4 (KB)",
        "Link Google Drive (Vídeo)", "Link Google Drive (Pôster)",
        "SHA-256", "Data/Hora Conclusão",
    ]
    assert REQUIRED_COLUMNS == esperado
    assert COLUMNS[:8] == esperado


def test_patch_parcial_nao_apaga_celulas():
    cells = RowUpdate(status=STATUS_DONE, sha256="abc").as_cells()
    assert cells == {"Status": STATUS_DONE, "SHA-256": "abc"}
    assert "Link Google Drive (Vídeo)" not in cells


def test_excel_tracker_round_trip():
    from tracker import ExcelTracker

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "c.xlsx"
        tracker = ExcelTracker(path)
        pares = list(prompts_data.iter_matrix(["A01"], None))
        tracker.bootstrap(pares)
        tracker.update("A01", "basal", RowUpdate(status=STATUS_DONE, mp4_kb=412.5, sha256="ff"))

        from openpyxl import load_workbook

        rows = list(load_workbook(path).active.iter_rows(values_only=True))
        assert rows[0][: len(COLUMNS)] == tuple(COLUMNS)
        assert len(rows) - 1 == 5
        linha = next(r for r in rows[1:] if r[1] == "basal")
        assert linha[2] == STATUS_DONE and linha[3] == 412.5 and linha[6] == "ff"

        # Reabrir não pode duplicar linhas.
        ExcelTracker(path).bootstrap(pares)
        assert len(list(load_workbook(path).active.iter_rows(values_only=True))) - 1 == 5


# --------------------------------------------------------------------------- #
# Cadeia FFmpeg (requer binários)
# --------------------------------------------------------------------------- #
def test_cadeia_ffmpeg_respeita_toda_a_especificacao():
    if not HAS_FFMPEG:  # pragma: no cover
        print("  (pulado: ffmpeg/ffprobe ausentes)")
        return

    from processor import probe, process_clip, synth_raw_clip

    with tempfile.TemporaryDirectory() as tmp:
        settings = load_settings({"work_dir": Path(tmp)})
        settings.ensure_dirs()
        raw = synth_raw_clip(settings, settings.raw_dir / "r.mp4", seed=1)
        result = process_clip(
            settings, raw, settings.out_dir / "basal.mp4", settings.out_dir / "basal.jpg"
        )

        assert result.width == 720 and result.height == 720      # 1:1
        assert result.fps == 24
        assert abs(result.duration_seconds - 3.5) < 0.05         # 2*2 - 0.5
        assert result.mp4_bytes <= 600 * 1024                    # teto de 600 KB
        assert result.jpg_bytes <= 60 * 1024                     # teto de 60 KB
        assert len(result.mp4_sha256) == 64 and len(result.jpg_sha256) == 64
        assert result.mp4_sha256 != result.jpg_sha256
        assert result.mp4_sha256 == sha256_file(Path(result.mp4_path))

        info = probe(settings, Path(result.mp4_path))
        assert info["width"] == 720 and info["height"] == 720

        import subprocess

        streams = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "stream=codec_type,profile,pix_fmt",
             "-of", "csv=p=0", result.mp4_path],
            capture_output=True, text=True,
        ).stdout
        assert "audio" not in streams, "o clipe final não pode conter faixa de áudio"
        assert "Main" in streams and "yuv420p" in streams


# --------------------------------------------------------------------------- #
# Runner standalone
# --------------------------------------------------------------------------- #
def main() -> int:
    testes = [
        (name, fn)
        for name, fn in sorted(globals().items())
        if name.startswith("test_") and callable(fn)
    ]
    falhas = 0
    print(f"Executando {len(testes)} testes…\n")
    for name, fn in testes:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            falhas += 1
            print(f"  ✗ {name}\n      {type(exc).__name__}: {exc}")
        else:
            print(f"  ✓ {name}")
    print(f"\n{len(testes) - falhas}/{len(testes)} passaram.")
    return 1 if falhas else 0


if __name__ == "__main__":
    sys.exit(main())
