"""Planilha de controle em tempo real — Excel (openpyxl) ou Google Sheets (gspread).

Cada par ator × estado ocupa exatamente uma linha, chaveada por ``(Ator, Estado)``.
O pipeline chama ``update()`` duas vezes por clipe: uma ao iniciar
(``⏳ Renderizando``) e outra ao concluir (``✅ Concluído``) ou falhar (``❌ Falha``).

O backend Excel grava em disco a cada atualização usando escrita atômica
(``tempfile`` + ``os.replace``), de modo que um Ctrl-C jamais deixa a planilha
corrompida — no pior caso ela perde a última linha.
"""

from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from config import Settings
from utils import LOG, retry, timestamp

# Ordem canônica das colunas — respeitada pelos dois backends.
# As oito primeiras são as obrigatórias da especificação; "Observação" é um
# extra opcional onde o motivo de uma falha é registrado, para que a coluna
# SHA-256 nunca carregue nada além de um hash.
COLUMNS: List[str] = [
    "Ator",
    "Estado",
    "Status",
    "Tamanho MP4 (KB)",
    "Link Google Drive (Vídeo)",
    "Link Google Drive (Pôster)",
    "SHA-256",
    "Data/Hora Conclusão",
    "Observação",
]
REQUIRED_COLUMNS: List[str] = COLUMNS[:8]

STATUS_PENDING = "⬜ Pendente"
STATUS_RUNNING = "⏳ Renderizando"
STATUS_DONE = "✅ Concluído"
STATUS_FAILED = "❌ Falha"
STATUS_SKIPPED = "↩️ Reaproveitado"


@dataclass
class RowUpdate:
    """Patch parcial de uma linha — campos ``None`` são preservados."""

    status: Optional[str] = None
    mp4_kb: Optional[float] = None
    video_link: Optional[str] = None
    poster_link: Optional[str] = None
    sha256: Optional[str] = None
    finished_at: Optional[str] = None
    note: Optional[str] = None

    def as_cells(self) -> Dict[str, object]:
        mapping = {
            "Status": self.status,
            "Tamanho MP4 (KB)": self.mp4_kb,
            "Link Google Drive (Vídeo)": self.video_link,
            "Link Google Drive (Pôster)": self.poster_link,
            "SHA-256": self.sha256,
            "Data/Hora Conclusão": self.finished_at,
            "Observação": self.note,
        }
        return {key: value for key, value in mapping.items() if value is not None}


# --------------------------------------------------------------------------- #
# Base
# --------------------------------------------------------------------------- #
class BaseTracker:
    """Contrato comum dos backends de planilha."""

    name = "base"

    def bootstrap(self, pairs: Sequence[Tuple[str, str]]) -> None:
        """Garante cabeçalho e uma linha por par ator × estado."""
        raise NotImplementedError

    def update(self, actor_id: str, state_id: str, patch: RowUpdate) -> None:
        raise NotImplementedError

    def describe(self) -> str:
        return self.name

    def close(self) -> None:  # pragma: no cover - opcional
        pass


class NullTracker(BaseTracker):
    """Backend ``none``: registra apenas no log."""

    name = "none"

    def bootstrap(self, pairs: Sequence[Tuple[str, str]]) -> None:
        LOG.info("Tracker desabilitado (%d clipes não serão registrados).", len(pairs))

    def update(self, actor_id: str, state_id: str, patch: RowUpdate) -> None:
        LOG.debug("Tracker(noop) %s/%s → %s", actor_id, state_id, patch.status)

    def describe(self) -> str:
        return "desabilitado"


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
class ExcelTracker(BaseTracker):
    """Mantém `controle_producao_ectoscopia.xlsx` atualizado a cada clipe."""

    name = "excel"

    HEADER_FILL = "FF1F3864"
    STATUS_FILLS = {
        STATUS_DONE: "FFD9EAD3",
        STATUS_RUNNING: "FFFFF2CC",
        STATUS_FAILED: "FFF4CCCC",
        STATUS_SKIPPED: "FFE6E6E6",
        STATUS_PENDING: "FFFFFFFF",
    }

    def __init__(self, path: Path) -> None:
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Instale `openpyxl` (pip install -r requirements.txt).") from exc

        self._Workbook = Workbook
        self._load_workbook = load_workbook
        self.path = Path(path)
        self._index: Dict[Tuple[str, str], int] = {}
        self._workbook = None
        self._sheet = None
        self._open()

    # -- ciclo de vida ------------------------------------------------- #
    def _open(self) -> None:
        if self.path.is_file():
            try:
                self._workbook = self._load_workbook(self.path)
                self._sheet = self._workbook.active
                self._reindex()
                LOG.info("Planilha existente carregada: %s (%d linhas).", self.path, len(self._index))
                return
            except Exception as exc:  # noqa: BLE001 - arquivo corrompido/aberto
                backup = self.path.with_suffix(".corrompido.xlsx")
                LOG.warning("Não foi possível abrir %s (%s). Movendo para %s.", self.path, exc, backup)
                try:
                    shutil.move(str(self.path), str(backup))
                except Exception:  # pragma: no cover
                    pass

        self._workbook = self._Workbook()
        self._sheet = self._workbook.active
        self._sheet.title = "Produção"
        self._sheet.append(COLUMNS)
        self._style_header()
        self._index = {}

    def _reindex(self) -> None:
        self._index = {}
        for row_idx, row in enumerate(self._sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row or row[0] is None:
                continue
            self._index[(str(row[0]), str(row[1]))] = row_idx

    # -- estilo --------------------------------------------------------- #
    def _style_header(self) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        widths = [8, 22, 18, 18, 46, 46, 68, 22, 52]
        for col_idx, (title, width) in enumerate(zip(COLUMNS, widths), start=1):
            cell = self._sheet.cell(row=1, column=col_idx, value=title)
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=self.HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            self._sheet.column_dimensions[get_column_letter(col_idx)].width = width
        self._sheet.freeze_panes = "A2"

    def _paint_status(self, row_idx: int, status: str) -> None:
        from openpyxl.styles import PatternFill

        color = self.STATUS_FILLS.get(status)
        if not color:
            return
        self._sheet.cell(row=row_idx, column=COLUMNS.index("Status") + 1).fill = PatternFill(
            "solid", fgColor=color
        )

    # -- API ------------------------------------------------------------ #
    def bootstrap(self, pairs: Sequence[Tuple[str, str]]) -> None:
        for actor_id, state_id in pairs:
            if (actor_id, state_id) in self._index:
                continue
            self._sheet.append(
                [actor_id, state_id, STATUS_PENDING] + [None] * (len(COLUMNS) - 3)
            )
            self._index[(actor_id, state_id)] = self._sheet.max_row
            self._paint_status(self._sheet.max_row, STATUS_PENDING)
        self._save()
        LOG.info("Planilha pronta: %s (%d linhas).", self.path, len(self._index))

    def update(self, actor_id: str, state_id: str, patch: RowUpdate) -> None:
        key = (actor_id, state_id)
        if key not in self._index:
            self.bootstrap([key])
        row_idx = self._index[key]

        for column, value in patch.as_cells().items():
            cell = self._sheet.cell(row=row_idx, column=COLUMNS.index(column) + 1, value=value)
            if column.startswith("Link") and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"
        if patch.status:
            self._paint_status(row_idx, patch.status)
        self._save()

    def _save(self) -> None:
        """Escrita atômica: nunca deixa um .xlsx meio-gravado no lugar do bom."""

        def _write() -> None:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            handle, tmp_name = tempfile.mkstemp(
                dir=str(self.path.parent or "."), prefix=".ctrl-", suffix=".xlsx"
            )
            os.close(handle)
            tmp_path = Path(tmp_name)
            try:
                self._workbook.save(tmp_path)
                os.replace(tmp_path, self.path)
            finally:
                if tmp_path.exists():
                    tmp_path.unlink(missing_ok=True)

        try:
            retry(_write, what=f"salvar {self.path.name}", attempts=4, base_delay=1.5, max_delay=10.0)
        except PermissionError:
            # Cenário clássico: a planilha está aberta no Excel e travada.
            fallback = self.path.with_name(f"{self.path.stem}.parcial.xlsx")
            LOG.warning("%s está bloqueada. Gravando em %s.", self.path, fallback)
            self._workbook.save(fallback)

    def describe(self) -> str:
        return f"Excel → {self.path}"


# --------------------------------------------------------------------------- #
# Google Sheets
# --------------------------------------------------------------------------- #
class SheetsTracker(BaseTracker):
    """Atualiza uma planilha do Google Sheets em tempo real via gspread."""

    name = "sheets"

    def __init__(self, settings: Settings) -> None:
        try:
            import gspread
            from google.oauth2.service_account import Credentials
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "Instale `gspread` e `google-auth` (pip install -r requirements.txt)."
            ) from exc

        creds_path = Path(settings.sheets_credentials_file)
        if not creds_path.is_file():
            raise RuntimeError(f"Credencial para Sheets não encontrada: {creds_path}")

        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        credentials = Credentials.from_service_account_file(str(creds_path), scopes=scopes)
        self._client = gspread.authorize(credentials)
        self._settings = settings
        self._spreadsheet = self._client.open_by_key(settings.sheets_spreadsheet_id)

        try:
            self._sheet = self._spreadsheet.worksheet(settings.sheets_worksheet)
        except Exception:
            self._sheet = self._spreadsheet.add_worksheet(
                title=settings.sheets_worksheet, rows=100, cols=len(COLUMNS)
            )
        self._index: Dict[Tuple[str, str], int] = {}
        LOG.info(
            "Google Sheets conectado: %s / %s", self._spreadsheet.title, settings.sheets_worksheet
        )

    def _a1(self, row_idx: int) -> str:
        from gspread.utils import rowcol_to_a1

        return f"{rowcol_to_a1(row_idx, 1)}:{rowcol_to_a1(row_idx, len(COLUMNS))}"

    def bootstrap(self, pairs: Sequence[Tuple[str, str]]) -> None:
        values = retry(self._sheet.get_all_values, what="Sheets: ler planilha", attempts=4)

        header_ok = bool(values) and values[0][: len(COLUMNS)] == COLUMNS
        if not header_ok:
            retry(
                lambda: self._sheet.update(range_name=self._a1(1), values=[COLUMNS]),
                what="Sheets: escrever cabeçalho",
                attempts=4,
            )
            retry(
                lambda: self._sheet.format(
                    self._a1(1),
                    {
                        "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                        "backgroundColor": {"red": 0.12, "green": 0.22, "blue": 0.39},
                        "horizontalAlignment": "CENTER",
                    },
                ),
                what="Sheets: formatar cabeçalho",
                attempts=3,
            )
            values = [COLUMNS]

        self._index = {}
        for row_idx, row in enumerate(values[1:], start=2):
            if len(row) >= 2 and row[0]:
                self._index[(row[0], row[1])] = row_idx

        missing = [pair for pair in pairs if pair not in self._index]
        if missing:
            next_row = max(self._index.values(), default=1) + 1
            block = [
                [actor_id, state_id, STATUS_PENDING] + [""] * (len(COLUMNS) - 3)
                for actor_id, state_id in missing
            ]
            retry(
                lambda: self._sheet.update(
                    range_name=f"A{next_row}", values=block, value_input_option="USER_ENTERED"
                ),
                what="Sheets: criar linhas",
                attempts=4,
            )
            for offset, pair in enumerate(missing):
                self._index[pair] = next_row + offset

        LOG.info("Google Sheets pronto (%d linhas).", len(self._index))

    def update(self, actor_id: str, state_id: str, patch: RowUpdate) -> None:
        from gspread.utils import rowcol_to_a1

        key = (actor_id, state_id)
        if key not in self._index:
            self.bootstrap([key])
        row_idx = self._index[key]

        requests = []
        for column, value in patch.as_cells().items():
            col_idx = COLUMNS.index(column) + 1
            cell_value = value
            if column.startswith("Link") and isinstance(value, str) and value.startswith("http"):
                label = "Vídeo" if "Vídeo" in column else "Pôster"
                cell_value = f'=HYPERLINK("{value}";"{label}")'
            requests.append(
                {"range": rowcol_to_a1(row_idx, col_idx), "values": [[cell_value]]}
            )

        if requests:
            retry(
                lambda: self._sheet.batch_update(requests, value_input_option="USER_ENTERED"),
                what=f"Sheets: atualizar {actor_id}/{state_id}",
                attempts=4,
            )

    def describe(self) -> str:
        return (
            f"Google Sheets → {self._spreadsheet.title} / {self._settings.sheets_worksheet}"
        )


# --------------------------------------------------------------------------- #
# Fábrica
# --------------------------------------------------------------------------- #
def build_tracker(settings: Settings) -> BaseTracker:
    """Instancia o backend configurado, degradando para Excel se o Sheets falhar."""
    backend = settings.tracker_backend
    if backend == "none":
        return NullTracker()
    if backend == "sheets":
        try:
            return SheetsTracker(settings)
        except Exception as exc:  # noqa: BLE001
            LOG.error(
                "Falha ao conectar no Google Sheets (%s). Caindo para o backend Excel.", exc
            )
            return ExcelTracker(Path(settings.excel_path))
    return ExcelTracker(Path(settings.excel_path))


def now(settings: Settings) -> str:
    return timestamp(settings.timezone)
